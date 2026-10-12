from functools import wraps
from datetime import datetime
from html import escape
from html.parser import HTMLParser
from io import BytesIO
import re
from pathlib import Path
import sqlite3
from zipfile import BadZipFile

from flask import Flask, flash, g, jsonify, redirect, render_template, request, send_file, session, url_for
from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from werkzeug.security import check_password_hash, generate_password_hash


BASE_DIR = Path(__file__).resolve().parent
DATABASE = BASE_DIR / "database.sqlite3"
DEFAULT_SECTIONS = (
    "Криптография",
    "Ошибки при отправке",
    "Ошибки при распаковке",
    "Ошибки на Hab",
)

app = Flask(__name__)
app.config["SECRET_KEY"] = "change-this-secret-key"


class RichTextSanitizer(HTMLParser):
    allowed_tags = {
        "p", "br", "strong", "b", "em", "i", "u", "s",
        "ol", "ul", "li", "blockquote", "h3", "h4", "a", "span",
    }

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.output = []

    def handle_starttag(self, tag, attrs):
        if tag not in self.allowed_tags:
            return
        safe_attrs = []
        for name, value in attrs:
            if tag == "a" and name == "href" and value:
                if value.startswith(("http://", "https://", "mailto:")):
                    safe_attrs.append(f'href="{escape(value, quote=True)}"')
                    safe_attrs.append('target="_blank"')
                    safe_attrs.append('rel="noopener noreferrer"')
            elif tag == "span" and name == "style" and value:
                styles = []
                for declaration in value.split(";"):
                    if ":" not in declaration:
                        continue
                    property_name, property_value = declaration.split(":", 1)
                    property_name = property_name.strip().lower()
                    property_value = property_value.strip()
                    if property_name in {"color", "background-color"} and re.fullmatch(
                        r"#[0-9a-fA-F]{3,8}|rgb\([^)]{1,30}\)|[a-zA-Z]{1,20}",
                        property_value,
                    ):
                        styles.append(f"{property_name}:{property_value}")
                if styles:
                    safe_attrs.append(f'style="{";".join(styles)}"')
        self.output.append(f"<{tag}{(' ' + ' '.join(safe_attrs)) if safe_attrs else ''}>")

    def handle_endtag(self, tag):
        if tag in self.allowed_tags and tag != "br":
            self.output.append(f"</{tag}>")

    def handle_data(self, data):
        self.output.append(escape(data))

    def handle_entityref(self, name):
        self.output.append(f"&{name};")

    def handle_charref(self, name):
        self.output.append(f"&#{name};")


def sanitize_rich_text(value):
    sanitizer = RichTextSanitizer()
    sanitizer.feed(value)
    sanitizer.close()
    return "".join(sanitizer.output).strip()


@app.template_filter("datetime_ru")
def format_datetime_ru(value):
    if not value:
        return ""
    try:
        parsed = datetime.fromisoformat(value)
        return parsed.strftime("%d.%m.%Y в %H:%M")
    except ValueError:
        return value


def get_db():
    if "db" not in g:
        g.db = sqlite3.connect(DATABASE)
        g.db.row_factory = sqlite3.Row
    return g.db


@app.teardown_appcontext
def close_db(_error=None):
    db = g.pop("db", None)
    if db is not None:
        db.close()


def init_db():
    db = get_db()
    db.executescript(
        """
        CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT NOT NULL UNIQUE,
            password_hash TEXT NOT NULL
        );

        CREATE TABLE IF NOT EXISTS cards (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            section TEXT NOT NULL DEFAULT 'Криптография',
            name TEXT NOT NULL,
            short_description TEXT NOT NULL,
            description TEXT NOT NULL,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        );

        CREATE TABLE IF NOT EXISTS card_proposals (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER NOT NULL,
            section TEXT NOT NULL,
            name TEXT NOT NULL,
            short_description TEXT NOT NULL,
            description TEXT NOT NULL,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (user_id) REFERENCES users (id)
        );

        CREATE TABLE IF NOT EXISTS user_favorites (
            user_id INTEGER NOT NULL,
            card_id INTEGER NOT NULL,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            PRIMARY KEY (user_id, card_id),
            FOREIGN KEY (user_id) REFERENCES users (id),
            FOREIGN KEY (card_id) REFERENCES cards (id)
        );

        CREATE TABLE IF NOT EXISTS news (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            title TEXT NOT NULL,
            short_description TEXT NOT NULL,
            content TEXT NOT NULL,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        );

        CREATE TABLE IF NOT EXISTS notes (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER NOT NULL,
            title TEXT NOT NULL,
            short_description TEXT NOT NULL,
            description TEXT NOT NULL,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (user_id) REFERENCES users (id)
        );

        CREATE TABLE IF NOT EXISTS card_comments (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            card_id INTEGER NOT NULL,
            user_id INTEGER NOT NULL,
            content TEXT NOT NULL,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            FOREIGN KEY (card_id) REFERENCES cards (id),
            FOREIGN KEY (user_id) REFERENCES users (id)
        );
        CREATE INDEX IF NOT EXISTS idx_card_comments_card_created
            ON card_comments (card_id, created_at, id);

        CREATE TABLE IF NOT EXISTS sections (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL UNIQUE COLLATE NOCASE
        );
        """
    )
    db.executemany(
        "INSERT OR IGNORE INTO sections (name) VALUES (?)",
        ((section,) for section in DEFAULT_SECTIONS),
    )
    columns = {
        row["name"] for row in db.execute("PRAGMA table_info(cards)").fetchall()
    }
    if "section" not in columns:
        db.execute(
            "ALTER TABLE cards ADD COLUMN section TEXT NOT NULL DEFAULT 'Криптография'"
        )
    user = db.execute("SELECT id FROM users WHERE username = ?", ("admin",)).fetchone()
    if user is None:
        db.execute(
            "INSERT INTO users (username, password_hash) VALUES (?, ?)",
            ("admin", generate_password_hash("admin123")),
        )
    db.commit()


def get_sections():
    return tuple(
        row["name"]
        for row in get_db().execute("SELECT name FROM sections ORDER BY id").fetchall()
    )


def login_required(view):
    @wraps(view)
    def wrapped_view(**kwargs):
        if "user_id" not in session:
            return redirect(url_for("login"))
        return view(**kwargs)

    return wrapped_view


def admin_required(view):
    @wraps(view)
    def wrapped_view(**kwargs):
        if session.get("username") != "admin":
            flash("Добавлять, редактировать и удалять карточки может только admin.", "error")
            return redirect(url_for("index"))
        return view(**kwargs)

    return wrapped_view


@app.context_processor
def inject_user_permissions():
    return {
        "is_admin": session.get("username") == "admin",
        "section_options": get_sections(),
    }


def validate_card_form():
    section = request.form.get("section", "").strip()
    name = request.form.get("name", "").strip()
    short_description = request.form.get("short_description", "").strip()
    description = sanitize_rich_text(request.form.get("description", ""))
    if section not in get_sections():
        flash("Выберите корректный раздел.", "error")
        return None
    if not name or not short_description or not description:
        flash("Заполните все обязательные поля.", "error")
        return None
    return section, name, short_description, description


def validate_section_form():
    name = request.form.get("name", "").strip()
    if not name:
        flash("Введите название раздела.", "error")
        return None
    if len(name) > 100:
        flash("Название раздела не должно превышать 100 символов.", "error")
        return None
    if get_db().execute(
        "SELECT 1 FROM sections WHERE name = ? COLLATE NOCASE", (name,)
    ).fetchone():
        flash("Раздел с таким названием уже существует.", "error")
        return None
    return name


def validate_proposal_form():
    return validate_card_form()


def validate_news_form():
    title = request.form.get("title", "").strip()
    short_description = request.form.get("short_description", "").strip()
    content = sanitize_rich_text(request.form.get("content", ""))
    if not title or not short_description or not content:
        flash("Заполните все обязательные поля новости.", "error")
        return None
    return title, short_description, content


def validate_note_form():
    title = request.form.get("title", "").strip()
    short_description = request.form.get("short_description", "").strip()
    description = sanitize_rich_text(request.form.get("description", ""))
    if not title or not short_description or not description:
        flash("Заполните все обязательные поля заметки.", "error")
        return None
    return title, short_description, description


@app.route("/login", methods=("GET", "POST"))
def login():
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")
        user = get_db().execute(
            "SELECT * FROM users WHERE username = ?", (username,)
        ).fetchone()
        if user is None or not check_password_hash(user["password_hash"], password):
            flash("Неверный логин или пароль.", "error")
        else:
            session.clear()
            session["user_id"] = user["id"]
            session["username"] = user["username"]
            return redirect(url_for("index"))
    return render_template("login.html")


@app.route("/register", methods=("GET", "POST"))
def register():
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")
        password_confirmation = request.form.get("password_confirmation", "")

        if len(username) < 3:
            flash("Логин должен содержать минимум 3 символа.", "error")
        elif len(password) < 6:
            flash("Пароль должен содержать минимум 6 символов.", "error")
        elif password != password_confirmation:
            flash("Пароли не совпадают.", "error")
        elif get_db().execute(
            "SELECT id FROM users WHERE username = ?", (username,)
        ).fetchone():
            flash("Пользователь с таким логином уже существует.", "error")
        else:
            db = get_db()
            cursor = db.execute(
                "INSERT INTO users (username, password_hash) VALUES (?, ?)",
                (username, generate_password_hash(password)),
            )
            db.commit()
            session.clear()
            session["user_id"] = cursor.lastrowid
            session["username"] = username
            return redirect(url_for("index"))
    return render_template("register.html")


@app.route("/notes")
@login_required
def notes():
    note_rows = get_db().execute(
        "SELECT * FROM notes WHERE user_id = ? ORDER BY updated_at DESC, id DESC",
        (session["user_id"],),
    ).fetchall()
    return render_template("notes.html", notes=note_rows)


@app.route("/notes/add", methods=("GET", "POST"))
@login_required
def add_note():
    if request.method == "POST":
        values = validate_note_form()
        if values is not None:
            db = get_db()
            db.execute(
                """
                INSERT INTO notes (user_id, title, short_description, description)
                VALUES (?, ?, ?, ?)
                """,
                (session["user_id"], *values),
            )
            db.commit()
            flash("Заметка добавлена.", "success")
            return redirect(url_for("notes"))
    return render_template("note_form.html", title="Добавить заметку", note=None)


@app.route("/notes/<int:note_id>/edit", methods=("GET", "POST"))
@login_required
def edit_note(note_id):
    db = get_db()
    note = db.execute(
        "SELECT * FROM notes WHERE id = ? AND user_id = ?",
        (note_id, session["user_id"]),
    ).fetchone()
    if note is None:
        flash("Заметка не найдена.", "error")
        return redirect(url_for("notes"))
    if request.method == "POST":
        values = validate_note_form()
        if values is not None:
            db.execute(
                """
                UPDATE notes
                SET title = ?, short_description = ?, description = ?,
                    updated_at = CURRENT_TIMESTAMP
                WHERE id = ? AND user_id = ?
                """,
                (*values, note_id, session["user_id"]),
            )
            db.commit()
            flash("Заметка обновлена.", "success")
            return redirect(url_for("notes"))
    return render_template("note_form.html", title="Редактировать заметку", note=note)


@app.post("/notes/<int:note_id>/delete")
@login_required
def delete_note(note_id):
    db = get_db()
    cursor = db.execute(
        "DELETE FROM notes WHERE id = ? AND user_id = ?",
        (note_id, session["user_id"]),
    )
    if cursor.rowcount == 0:
        flash("Заметка не найдена.", "error")
    else:
        db.commit()
        flash("Заметка удалена.", "success")
    return redirect(url_for("notes"))


@app.route("/news")
@login_required
def news():
    news_items = get_db().execute(
        "SELECT * FROM news ORDER BY created_at DESC, id DESC"
    ).fetchall()
    return render_template("news.html", news_items=news_items)


@app.route("/news/add", methods=("GET", "POST"))
@login_required
@admin_required
def add_news():
    if request.method == "POST":
        values = validate_news_form()
        if values is not None:
            get_db().execute(
                """
                INSERT INTO news (title, short_description, content)
                VALUES (?, ?, ?)
                """,
                values,
            )
            get_db().commit()
            flash("Новость добавлена.", "success")
            return redirect(url_for("news"))
    return render_template("news_form.html", title="Добавить новость")


@app.route("/news/<int:news_id>/edit", methods=("GET", "POST"))
@login_required
@admin_required
def edit_news(news_id):
    db = get_db()
    item = db.execute("SELECT * FROM news WHERE id = ?", (news_id,)).fetchone()
    if item is None:
        flash("Новость не найдена.", "error")
        return redirect(url_for("news"))
    if request.method == "POST":
        values = validate_news_form()
        if values is not None:
            db.execute(
                """
                UPDATE news
                SET title = ?, short_description = ?, content = ?,
                    updated_at = CURRENT_TIMESTAMP
                WHERE id = ?
                """,
                (*values, news_id),
            )
            db.commit()
            flash("Новость обновлена.", "success")
            return redirect(url_for("news"))
    return render_template("news_form.html", title="Редактировать новость", news_item=item)


@app.route("/proposals/add", methods=("GET", "POST"))
@login_required
def add_proposal():
    if request.method == "POST":
        values = validate_proposal_form()
        if values is not None:
            db = get_db()
            db.execute(
                """
                INSERT INTO card_proposals
                    (user_id, section, name, short_description, description)
                VALUES (?, ?, ?, ?, ?)
                """,
                (session["user_id"], *values),
            )
            db.commit()
            flash("Ваша карточка направлена на рассмотрение.", "success")
            return redirect(url_for("index"))
    return render_template(
        "card_form.html",
        card=None,
        title="Предложить карточку",
        sections=get_sections(),
        form_action=url_for("add_proposal"),
        form_submit_label="Предложить",
    )


@app.post("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


@app.route("/")
@login_required
def index():
    db = get_db()
    query = request.args.get("q", "").strip()
    card_id = request.args.get("card_id", "").strip()
    if query or card_id:
        conditions = []
        params = []
        if query:
            conditions.append("(name LIKE ? OR short_description LIKE ?)")
            search_value = f"%{query}%"
            params.extend((search_value, search_value))
        if card_id:
            if not card_id.isdigit():
                flash("Номер карточки должен содержать только цифры.", "error")
                return redirect(url_for("index"))
            conditions.append("id = ?")
            params.append(int(card_id))
        cards = db.execute(
            f"SELECT * FROM cards WHERE {' AND '.join(conditions)} "
            "ORDER BY updated_at DESC, id DESC",
            params,
        ).fetchall()
        return render_template(
            "index.html",
            cards=cards,
            search_active=True,
            search_query=query,
            search_card_id=card_id,
            favorite_ids=get_favorite_ids(),
        )
    sections = [
        {
            "name": section,
            "count": db.execute(
                "SELECT COUNT(*) AS count FROM cards WHERE section = ?", (section,)
            ).fetchone()["count"],
        }
        for section in get_sections()
    ]
    return render_template("index.html", sections=sections)


@app.route("/sections/add", methods=("GET", "POST"))
@login_required
@admin_required
def add_section():
    if request.method == "POST":
        name = validate_section_form()
        if name is not None:
            try:
                get_db().execute("INSERT INTO sections (name) VALUES (?)", (name,))
                get_db().commit()
            except sqlite3.IntegrityError:
                flash("Раздел с таким названием уже существует.", "error")
            else:
                flash("Раздел добавлен.", "success")
                return redirect(url_for("index"))
    return render_template("section_form.html", title="Добавить раздел")


@app.route("/cards/import", methods=("GET", "POST"))
@login_required
@admin_required
def import_cards():
    if request.method == "POST":
        upload = request.files.get("file")
        errors = []
        if request.content_length and request.content_length > 10 * 1024 * 1024:
            errors.append("Файл не должен превышать 10 МБ.")
        elif upload is None or not upload.filename:
            errors.append("Выберите файл Excel для загрузки.")
        elif Path(upload.filename).suffix.lower() != ".xlsx":
            errors.append("Поддерживаются только файлы формата .xlsx.")
        if errors:
            return render_template("card_import.html", errors=errors)

        workbook = None
        try:
            workbook = load_workbook(upload.stream, read_only=True, data_only=True)
            worksheet = workbook.active
            if worksheet is None:
                errors.append("В книге Excel не найден лист с данными.")
            else:
                header_values = next(
                    worksheet.iter_rows(min_row=1, max_row=1, values_only=True),
                    (),
                )
                headers = [
                    str(value).strip().casefold() if value is not None else ""
                    for value in header_values
                ]
                required_headers = (
                    "раздел",
                    "наименование",
                    "краткое описание",
                    "описание",
                )
                missing_headers = [header for header in required_headers if header not in headers]
                duplicate_headers = [
                    header for header in required_headers if headers.count(header) > 1
                ]
                if missing_headers or duplicate_headers:
                    if missing_headers:
                        errors.append(
                            "В первой строке отсутствуют столбцы: "
                            + ", ".join(missing_headers)
                            + "."
                        )
                    if duplicate_headers:
                        errors.append(
                            "В первой строке повторяются столбцы: "
                            + ", ".join(duplicate_headers)
                            + "."
                        )
                else:
                    indexes = {header: headers.index(header) for header in required_headers}
                    valid_sections = set(get_sections())
                    valid_rows = []
                    data_row_count = 0
                    for row_number, row in enumerate(
                        worksheet.iter_rows(min_row=2, values_only=True), start=2
                    ):
                        if not any(value is not None and str(value).strip() for value in row):
                            continue
                        data_row_count += 1
                        if data_row_count > 5000:
                            errors.append("В файле больше 5000 строк с данными.")
                            break

                        values = {
                            header: (
                                str(row[indexes[header]]).strip()
                                if indexes[header] < len(row)
                                and row[indexes[header]] is not None
                                else ""
                            )
                            for header in required_headers
                        }
                        row_errors = []
                        if not values["раздел"]:
                            row_errors.append("не указан раздел")
                        elif values["раздел"] not in valid_sections:
                            row_errors.append(
                                f"раздел «{values['раздел']}» не существует"
                            )
                        if not values["наименование"]:
                            row_errors.append("не заполнено наименование")
                        elif len(values["наименование"]) > 200:
                            row_errors.append("наименование длиннее 200 символов")
                        if not values["краткое описание"]:
                            row_errors.append("не заполнено краткое описание")
                        elif len(values["краткое описание"]) > 300:
                            row_errors.append("краткое описание длиннее 300 символов")
                        description = sanitize_rich_text(values["описание"])
                        if not description:
                            row_errors.append("не заполнено описание")
                        if row_errors:
                            errors.append(
                                f"Строка {row_number}: " + "; ".join(row_errors) + "."
                            )
                        else:
                            valid_rows.append(
                                (
                                    values["раздел"],
                                    values["наименование"],
                                    values["краткое описание"],
                                    description,
                                )
                            )
        except (InvalidFileException, BadZipFile, OSError, ValueError):
            errors.append("Не удалось прочитать файл. Проверьте, что это корректная книга .xlsx.")
        finally:
            if workbook is not None:
                workbook.close()

        if errors:
            displayed_errors = errors[:100]
            if len(errors) > len(displayed_errors):
                displayed_errors.append(
                    f"И еще ошибок: {len(errors) - len(displayed_errors)}."
                )
            return render_template("card_import.html", errors=displayed_errors)

        if not valid_rows:
            return render_template(
                "card_import.html",
                errors=["В файле нет строк с карточками для импорта."],
            )

        db = get_db()
        try:
            db.execute("BEGIN")
            db.executemany(
                """
                INSERT INTO cards (section, name, short_description, description)
                VALUES (?, ?, ?, ?)
                """,
                valid_rows,
            )
            db.commit()
        except sqlite3.Error:
            db.rollback()
            app.logger.exception("Ошибка массового импорта карточек из XLSX.")
            return render_template(
                "card_import.html",
                errors=["Не удалось сохранить карточки. Данные из файла не были импортированы."],
            )

        flash(f"Добавлено карточек: {len(valid_rows)}.", "success")
        return redirect(url_for("index"))

    return render_template("card_import.html")


@app.get("/cards/import/template")
@login_required
@admin_required
def download_card_import_template():
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Карточки"
    worksheet.append(("Раздел", "Наименование", "Краткое описание", "Описание"))
    worksheet.freeze_panes = "A2"
    for column, width in (("A", 30), ("B", 35), ("C", 45), ("D", 70)):
        worksheet.column_dimensions[column].width = width

    output = BytesIO()
    workbook.save(output)
    workbook.close()
    output.seek(0)
    return send_file(
        output,
        as_attachment=True,
        download_name="шаблон_карточек.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.route("/sections/<path:section>")
@login_required
def section_cards(section):
    if section not in get_sections():
        flash("Раздел не найден.", "error")
        return redirect(url_for("index"))
    query = request.args.get("q", "").strip()
    card_id = request.args.get("card_id", "").strip()
    conditions = ["section = ?"]
    params = [section]
    if query:
        conditions.append("(name LIKE ? OR short_description LIKE ?)")
        search_value = f"%{query}%"
        params.extend((search_value, search_value))
    if card_id:
        if not card_id.isdigit():
            flash("Номер карточки должен содержать только цифры.", "error")
            return redirect(url_for("section_cards", section=section))
        conditions.append("id = ?")
        params.append(int(card_id))
    cards = get_db().execute(
        f"SELECT * FROM cards WHERE {' AND '.join(conditions)} "
        "ORDER BY updated_at DESC, id DESC",
        params,
    ).fetchall()
    return render_template(
        "index.html",
        cards=cards,
        section=section,
        search_active=bool(query or card_id),
        section_search_active=bool(query or card_id),
        section_search_query=query,
        section_search_card_id=card_id,
        favorite_ids=get_favorite_ids(),
    )


@app.route("/cards/<int:card_id>/comments", methods=("GET", "POST"))
@login_required
def card_comments(card_id):
    db = get_db()
    if db.execute("SELECT id FROM cards WHERE id = ?", (card_id,)).fetchone() is None:
        return jsonify(error="Карточка не найдена."), 404

    if request.method == "POST":
        content = request.form.get("content", "").strip()
        if not content:
            return jsonify(error="Введите текст комментария."), 400
        if len(content) > 2000:
            return jsonify(error="Комментарий не должен превышать 2000 символов."), 400
        cursor = db.execute(
            """
            INSERT INTO card_comments (card_id, user_id, content)
            VALUES (?, ?, ?)
            """,
            (card_id, session["user_id"], content),
        )
        db.commit()
        comment = db.execute(
            """
            SELECT card_comments.id, card_comments.content,
                   card_comments.created_at, users.username
            FROM card_comments
            JOIN users ON users.id = card_comments.user_id
            WHERE card_comments.id = ?
            """,
            (cursor.lastrowid,),
        ).fetchone()
        return jsonify(
            id=comment["id"],
            username=comment["username"],
            content=comment["content"],
            created_at=comment["created_at"],
        ), 201

    comments = db.execute(
        """
        SELECT card_comments.id, card_comments.content,
               card_comments.created_at, users.username
        FROM card_comments
        JOIN users ON users.id = card_comments.user_id
        WHERE card_comments.card_id = ?
        ORDER BY card_comments.created_at ASC, card_comments.id ASC
        """,
        (card_id,),
    ).fetchall()
    return jsonify(comments=[
        {
            "id": comment["id"],
            "username": comment["username"],
            "content": comment["content"],
            "created_at": comment["created_at"],
        }
        for comment in comments
    ])


def get_favorite_ids():
    if session.get("username") == "admin":
        return set()
    rows = get_db().execute(
        "SELECT card_id FROM user_favorites WHERE user_id = ?", (session["user_id"],)
    ).fetchall()
    return {row["card_id"] for row in rows}


@app.post("/cards/<int:card_id>/favorite")
@login_required
def toggle_favorite(card_id):
    is_ajax_request = request.headers.get("X-Requested-With") == "XMLHttpRequest"
    if session.get("username") == "admin":
        if is_ajax_request:
            return jsonify(error="Избранное доступно только обычным пользователям."), 403
        flash("Избранное доступно только обычным пользователям.", "error")
        return redirect(request.referrer or url_for("index"))
    db = get_db()
    card = db.execute("SELECT id FROM cards WHERE id = ?", (card_id,)).fetchone()
    if card is None:
        if is_ajax_request:
            return jsonify(error="Карточка не найдена."), 404
        flash("Карточка не найдена.", "error")
        return redirect(request.referrer or url_for("index"))
    favorite = db.execute(
        "SELECT 1 FROM user_favorites WHERE user_id = ? AND card_id = ?",
        (session["user_id"], card_id),
    ).fetchone()
    if favorite:
        db.execute(
            "DELETE FROM user_favorites WHERE user_id = ? AND card_id = ?",
            (session["user_id"], card_id),
        )
    else:
        db.execute(
            "INSERT INTO user_favorites (user_id, card_id) VALUES (?, ?)",
            (session["user_id"], card_id),
        )
    db.commit()
    if is_ajax_request:
        return jsonify(is_favorite=not bool(favorite))
    return redirect(request.referrer or url_for("index"))


@app.route("/favorites")
@login_required
def favorites():
    if session.get("username") == "admin":
        flash("Избранное доступно только обычным пользователям.", "error")
        return redirect(url_for("index"))
    cards = get_db().execute(
        """
        SELECT cards.*
        FROM cards
        JOIN user_favorites
          ON user_favorites.card_id = cards.id
        WHERE user_favorites.user_id = ?
        ORDER BY user_favorites.created_at DESC, cards.id DESC
        """,
        (session["user_id"],),
    ).fetchall()
    return render_template(
        "favorites.html", cards=cards, favorite_ids=get_favorite_ids()
    )


@app.route("/cards/add", methods=("GET", "POST"))
@login_required
@admin_required
def add_card():
    if request.method == "POST":
        values = validate_card_form()
        if values is not None:
            get_db().execute(
                """
                INSERT INTO cards (section, name, short_description, description)
                VALUES (?, ?, ?, ?)
                """,
                values,
            )
            get_db().commit()
            flash("Карточка добавлена.", "success")
            return redirect(url_for("section_cards", section=values[0]))
    return render_template(
        "card_form.html", card=None, title="Добавить карточку", sections=get_sections()
    )


@app.route("/cards/<int:card_id>/edit", methods=("GET", "POST"))
@login_required
@admin_required
def edit_card(card_id):
    db = get_db()
    card = db.execute("SELECT * FROM cards WHERE id = ?", (card_id,)).fetchone()
    if card is None:
        flash("Карточка не найдена.", "error")
        return redirect(url_for("index"))

    if request.method == "POST":
        values = validate_card_form()
        if values is not None:
            db.execute(
                """
                UPDATE cards
                SET section = ?, name = ?, short_description = ?, description = ?,
                    updated_at = CURRENT_TIMESTAMP
                WHERE id = ?
                """,
                (*values, card_id),
            )
            db.commit()
            flash("Карточка обновлена.", "success")
            return redirect(url_for("section_cards", section=values[0]))
    return render_template(
        "card_form.html", card=card, title="Редактировать карточку", sections=get_sections()
    )


@app.post("/cards/<int:card_id>/delete")
@login_required
@admin_required
def delete_card(card_id):
    db = get_db()
    card = db.execute("SELECT section FROM cards WHERE id = ?", (card_id,)).fetchone()
    if card is None:
        flash("Карточка не найдена.", "error")
        return redirect(url_for("index"))
    db.execute("DELETE FROM card_comments WHERE card_id = ?", (card_id,))
    db.execute("DELETE FROM user_favorites WHERE card_id = ?", (card_id,))
    db.execute("DELETE FROM cards WHERE id = ?", (card_id,))
    db.commit()
    flash("Карточка удалена.", "success")
    return redirect(url_for("section_cards", section=card["section"]))


@app.route("/proposals")
@login_required
@admin_required
def proposals():
    proposal_rows = get_db().execute(
        """
        SELECT card_proposals.*, users.username AS submitted_by
        FROM card_proposals
        JOIN users ON users.id = card_proposals.user_id
        ORDER BY card_proposals.created_at ASC, card_proposals.id ASC
        """
    ).fetchall()
    return render_template("proposals.html", proposals=proposal_rows)


@app.route("/proposals/<int:proposal_id>/edit", methods=("GET", "POST"))
@login_required
@admin_required
def edit_proposal(proposal_id):
    db = get_db()
    proposal = db.execute(
        "SELECT * FROM card_proposals WHERE id = ?", (proposal_id,)
    ).fetchone()
    if proposal is None:
        flash("Предложение не найдено.", "error")
        return redirect(url_for("proposals"))
    if request.method == "POST":
        values = validate_proposal_form()
        if values is not None:
            db.execute(
                """
                UPDATE card_proposals
                SET section = ?, name = ?, short_description = ?,
                    description = ?, updated_at = CURRENT_TIMESTAMP
                WHERE id = ?
                """,
                (*values, proposal_id),
            )
            db.commit()
            flash("Предложение обновлено.", "success")
            return redirect(url_for("proposals"))
    return render_template(
        "card_form.html",
        card=proposal,
        title="Редактировать предложение",
        sections=get_sections(),
        form_action=url_for("edit_proposal", proposal_id=proposal_id),
        form_submit_label="Сохранить",
    )


@app.post("/proposals/<int:proposal_id>/accept")
@login_required
@admin_required
def accept_proposal(proposal_id):
    db = get_db()
    proposal = db.execute(
        "SELECT * FROM card_proposals WHERE id = ?", (proposal_id,)
    ).fetchone()
    if proposal is None:
        flash("Предложение не найдено.", "error")
        return redirect(url_for("proposals"))
    db.execute(
        """
        INSERT INTO cards (section, name, short_description, description)
        VALUES (?, ?, ?, ?)
        """,
        (proposal["section"], proposal["name"], proposal["short_description"], proposal["description"]),
    )
    db.execute("DELETE FROM card_proposals WHERE id = ?", (proposal_id,))
    db.commit()
    flash("Предложение принято и добавлено в базу.", "success")
    return redirect(url_for("proposals"))


@app.post("/proposals/<int:proposal_id>/reject")
@login_required
@admin_required
def reject_proposal(proposal_id):
    db = get_db()
    proposal = db.execute(
        "SELECT id FROM card_proposals WHERE id = ?", (proposal_id,)
    ).fetchone()
    if proposal is None:
        flash("Предложение не найдено.", "error")
        return redirect(url_for("proposals"))
    db.execute("DELETE FROM card_proposals WHERE id = ?", (proposal_id,))
    db.commit()
    flash("Предложение отклонено.", "success")
    return redirect(url_for("proposals"))


with app.app_context():
    init_db()


if __name__ == "__main__":
    app.run(debug=True)
